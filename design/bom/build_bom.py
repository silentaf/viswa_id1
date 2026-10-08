"""VisionAid Stage 2: prototype BOM + power budget workbook (VisionAid_BOM_Power.xlsx).
Prices: 'Quote' = listing checked on 8 Oct 2026 (source given); 'Estimate' = typical Indian retail,
must be verified before ordering. Power figures are design estimates (datasheet / assumption),
NOT measurements - they are verified in Stage 3 with a USB-C power meter.
Run: python3 design/bom/build_bom.py
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))

BOM = [
    # group, item, part / spec, qty, unit INR, basis, source, origin, lead time
    ('Compute', 'AI computer', 'Raspberry Pi 5, 4 GB', 1, 12000, 'Quote', 'KSP Electronics listing (in stock); guides quote 5.5-7.5k', 'Imported, local distributor', '2-4 days'),
    ('Compute', 'Active cooler', 'Raspberry Pi 5 Active Cooler', 1, 550, 'Estimate', 'Thingbits / Robocraze', 'Imported, local distributor', '2-4 days'),
    ('Compute', 'Storage', '64 GB microSD, A2', 1, 700, 'Estimate', 'Local / Amazon.in', 'Local', '1-3 days'),
    ('Compute', 'FPGA safety island', 'Sipeed Tang Nano 9K (GW1NR-9)', 1, 2149, 'Quote', 'Probots.co.in (incl. GST)', 'Imported, local distributor', '1-3 days'),
    ('Sensing', 'Camera', 'Raspberry Pi Camera Module 3 (standard, autofocus)', 1, 3250, 'Quote', 'Thingbits.in', 'Imported, local distributor', '2-4 days'),
    ('Sensing', 'Camera cable', 'Pi 5 camera FFC 22-15 pin, 300 mm', 1, 200, 'Estimate', 'Robu / Thingbits', 'Imported, local distributor', '2-4 days'),
    ('Sensing', 'Ultrasonic sensor', 'RCWL-1601 (3.3 V, HC-SR04 pin-out), incl. 1 spare', 4, 150, 'Estimate', 'Robu / Probots', 'Imported, local distributor', '2-4 days'),
    ('Sensing', 'Drop-off LiDAR', 'Benewake TF-Luna (0.2-8 m, UART)', 1, 2118, 'Quote', 'Probots.co.in (incl. GST)', 'Imported, local distributor', '1-3 days'),
    ('Sensing', 'IMU', 'MPU-6050 (GY-521) module', 1, 200, 'Estimate', 'Robu', 'Local', '2-4 days'),
    ('Feedback', 'Haptic motors', '10 mm coin ERM, 3 V', 2, 60, 'Estimate', 'Robu', 'Local', '2-4 days'),
    ('Feedback', 'Audio', 'Bluetooth bone-conduction headset (budget)', 1, 2000, 'Estimate', 'Amazon.in', 'Imported', '2-5 days'),
    ('Power', 'Battery', '10,000 mAh USB-C power bank, BIS-certified', 1, 1200, 'Estimate', 'Amazon.in / local', 'Local brand', '1-3 days'),
    ('Carrier PCB', 'PCB fabrication', '2-layer FR-4 85 x 56 mm, 5 pcs (this design)', 1, 2000, 'Estimate', 'LionCircuits (Bengaluru)', 'Domestic fab', '5-7 days'),
    ('Carrier PCB', 'SMD parts', 'AMS1117, AO3400A x3, SS14, 1N4148W x3, R/C/LED, PTC (see KiCad BOM)', 1, 500, 'Estimate', 'Robu / Probots', 'Local', '2-4 days'),
    ('Carrier PCB', 'Connectors', 'JST-XH sockets + housings + crimps, 2x20 header, 2x 1x24 sockets, buzzer', 1, 600, 'Estimate', 'Robu', 'Local', '2-4 days'),
    ('Carrier PCB', 'Ribbon', '40-way IDC ribbon, 2 sockets, ~10 cm', 1, 150, 'Estimate', 'Robu', 'Local', '2-4 days'),
    ('Mechanical', 'Filament', 'PETG 1 kg (enclosure ~140 g solid volume)', 1, 1200, 'Estimate', 'Robu / Amazon.in', 'Local', '2-4 days'),
    ('Mechanical', 'Fasteners', 'M2.5 x 20 mm standoffs x4, M3 heat-set inserts, M2.5/M3 screws', 1, 400, 'Estimate', 'Robu', 'Local', '2-4 days'),
    ('Mechanical', 'Chest harness', 'GoPro-style chest harness + M5 thumb screw', 1, 600, 'Estimate', 'Amazon.in', 'Imported', '2-5 days'),
    ('Mechanical', 'Buttons', '3 x 12 mm tactile buttons with caps (READ / MODE / QUIET)', 1, 150, 'Estimate', 'Robu', 'Local', '2-4 days'),
]

POWER = [
    # block, rail, current mA, voltage, W, basis
    ('Raspberry Pi 5 (camera + YOLO running)', '5 V', None, 5.0, 6.0, 'Assumption - to be measured (Stage 3, USB-C meter)'),
    ('Camera Module 3', 'Pi', None, 3.3, 0.25, 'Assumption'),
    ('Tang Nano 9K (FPGA + on-board USB bridge)', '5 V', 80, 5.0, None, 'Assumption'),
    ('TF-Luna LiDAR', '5 V', 70, 5.0, None, 'Datasheet: average current <= 70 mA'),
    ('3 x RCWL-1601 ultrasonic', '3.3 V', 9, 3.3, None, 'Assumption (~3 mA each)'),
    ('2 x coin motors (30 % activity)', '3.3 V', 50, 3.3, None, 'Assumption (~80 mA each when on)'),
    ('MPU-6050 IMU', '3.3 V', 4, 3.3, None, 'Datasheet-level estimate'),
    ('AMS1117 loss (5 V -> 3.3 V, ~65 mA avg)', '5 V', None, None, 0.11, '(5 - 3.3) V x 65 mA'),
    ('Buzzer, LEDs (rare)', '5 V', 5, 5.0, None, 'Assumption'),
]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Prototype BOM'
    head = ['Group', 'Item', 'Part / specification', 'Qty', 'Unit price (INR)', 'Line total (INR)',
            'Price basis', 'Source', 'Origin', 'Lead time']
    ws.append(['VisionAid prototype BOM (Stage 2 design). Quote = listing checked 8 Oct 2026; Estimate = verify before ordering.'])
    ws.append(head)
    for r, row in enumerate(BOM, start=3):
        g, item, spec, q, u, basis, src, origin, lead = row
        ws.append([g, item, spec, q, u, '=D%d*E%d' % (r, r), basis, src, origin, lead])
    last = 2 + len(BOM)
    ws.append([])
    ws.append(['', 'TOTAL', '', '', '', '=SUM(F3:F%d)' % last])
    ws.append(['', 'Quoted items subtotal', '', '', '', '=SUMIF(G3:G%d,"Quote",F3:F%d)' % (last, last)])
    ws.append(['', 'Grant', '', '', '', 30000])
    ws.append(['', 'Reserve (grant - total)', '', '', '', '=F%d-F%d' % (last + 4, last + 2)])
    ws.append([])
    ws.append(['', 'BUDGET RULE', 'At the highest Pi 5 quote (INR 12,000) the build is ~INR 700 over the grant. '
               'At the INR 5,500-7,500 distributor price it is ~INR 24-26k. Rule: buy the Pi 5 4 GB only at '
               '<= INR 9,000, otherwise a Pi 4 4 GB (lower AI frame rate, same design).'])
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True, color='C0392B')
    for c in range(1, 11):
        ws.cell(row=2, column=c).font = Font(bold=True, color='FFFFFF')
        ws.cell(row=2, column=c).fill = PatternFill('solid', fgColor='1F5FA8')
    ws['A1'].font = Font(italic=True)
    for r in (last + 2, last + 3, last + 4, last + 5):
        ws.cell(row=r, column=2).font = Font(bold=True)
        ws.cell(row=r, column=6).font = Font(bold=True)
    widths = [12, 20, 62, 5, 16, 16, 12, 44, 26, 11]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ps = wb.create_sheet('Power budget')
    ps.append(['VisionAid power budget - DESIGN ESTIMATE (not measured). Verify in Stage 3 with a USB-C power meter.'])
    ps.append(['Block', 'Rail', 'Current (mA)', 'Voltage (V)', 'Power (W)', 'Basis'])
    for r, (blk, rail, ma, v, w, basis) in enumerate(POWER, start=3):
        pw = w if w is not None else '=C%d*D%d/1000' % (r, r)
        ps.append([blk, rail, ma, v, pw, basis])
    lastp = 2 + len(POWER)
    ps.append([])
    ps.append(['Total average power (W)', '', '', '', '=SUM(E3:E%d)' % lastp])
    ps.append(['Power bank energy (Wh): 10,000 mAh x 3.7 V', '', '', '', 37])
    ps.append(['Boost / cable efficiency (assumption)', '', '', '', 0.85])
    ps.append(['Estimated runtime (h)', '', '', '', '=E%d*E%d/E%d' % (lastp + 3, lastp + 4, lastp + 2)])
    ps.append(['Target', '', '', '', '>= 4 h'])
    for c in range(1, 7):
        ps.cell(row=2, column=c).font = Font(bold=True, color='FFFFFF')
        ps.cell(row=2, column=c).fill = PatternFill('solid', fgColor='1F5FA8')
    for i, w in enumerate([46, 8, 13, 12, 11, 52], start=1):
        ps.column_dimensions[get_column_letter(i)].width = w
    wb.save(os.path.join(HERE, 'VisionAid_BOM_Power.xlsx'))

    total = sum(q * u for _, _, _, q, u, *_ in BOM)
    quoted = sum(q * u for _, _, _, q, u, b, *_ in BOM if b == 'Quote')
    pw = 0.0
    for blk, rail, ma, v, w, basis in POWER:
        pw += w if w is not None else ma * v / 1000
    print('BOM total INR %d (quoted part %d), reserve %d' % (total, quoted, 30000 - total))
    print('power %.2f W -> runtime %.1f h on 10,000 mAh' % (pw, 37 * 0.85 / pw))


if __name__ == '__main__':
    main()
